import json
import openpyxl
import os
import sys
import re
import yaml

# Function to identify cells with formulas
def is_formula(cell):
    return isinstance(cell, openpyxl.cell.cell.Cell) and cell.data_type == 'f'

def process_workbook(file_name, output_file_json, output_file_yaml):
    # Load workbook
    wb = openpyxl.load_workbook(file_name)

    # Convert named ranges to their cell coordinate equivalents
    named_ranges = {name: wb.defined_names[name].destinations for name in wb.defined_names}

    cell_to_names = {}
    for name, cells in named_ranges.items():
        for cell_range in cells:
            sheet, cell_coord = cell_range
            cell_coord_clean = cell_coord.replace('$', '')
            key = f'{sheet}:{cell_coord_clean}'
            cell_to_names[key] = name

    named_formulas = {}
    sheet_formulas = {}

    for sheet in wb.sheetnames:
        cell_formulas = {}

        for row in wb[sheet].iter_rows():
            for cell in row:
                formula = cell.value
                if is_formula(cell):
                    formula = formula.replace('_xlfn.', '')
                    for coord, named_range in cell_to_names.items():
                        formula = re.sub(r'[A-Za-z0-9_\\\.]+(?: [A-Za-z0-9_\\\.]*)*!' + coord, named_range, formula.replace('$', ''))
                        formula = formula.replace(coord, named_range)
                if formula:
                    cell_formulas[cell.coordinate] = formula
                    cell_coord_clean = cell.coordinate.replace('$', '')
                    key = f'{sheet}:{cell_coord_clean}'
                    if key in cell_to_names.keys():
                        named_formulas[cell_to_names[key]] = formula
        sheet_formulas[sheet] = cell_formulas

    sheet_formulas["Names"] = named_formulas

    with open(output_file_json, 'w') as f:
        f.write(json.dumps(sheet_formulas, indent=1, default=str, ensure_ascii=False))

    with open(output_file_yaml, 'w') as f:
        yaml.dump(sheet_formulas, f, default_flow_style=False)

def main():
    input_file = sys.argv[1] if len(sys.argv) > 1 else None
    output_file_json = sys.argv[2] if len(sys.argv) > 2 else "output.json"
    output_file_yaml = sys.argv[3] if len(sys.argv) > 3 else "output.yaml"

    if input_file:
        process_workbook(input_file, output_file_json, output_file_yaml)
    else:
        for file in os.listdir():
            if file.endswith(".xlsx"):
                process_workbook(file, file.replace(".xlsx", ".json"), file.replace(".xlsx", ".yaml"))

if __name__ == "__main__":
    main()
